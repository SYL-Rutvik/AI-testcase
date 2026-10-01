import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.left_margin = Inches(1.2)
section.right_margin = Inches(1.0)
section.top_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)

# Header
header = section.header
p_head = header.paragraphs[0]
p_head.text = "24SOEIT13019\t\tCE738 - 7th Semester Project Report"

# Footer
footer = section.footer
p_foot = footer.paragraphs[0]
run_left = p_foot.add_run("CE/IT\t\tPage ")
fld = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
p_foot._p.append(fld)

doc.add_paragraph("Test Document for Guidelines.")
doc.save("test_page_no.docx")
print("Saved test_page_no.docx successfully!")
